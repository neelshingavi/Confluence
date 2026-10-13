"""End-to-end API smoke test using disposable SQLite data (not a PostgreSQL concurrency test)."""
from __future__ import annotations

import base64
import io
import json
import os
import secrets
import uuid
import zipfile
from pathlib import Path

# The API module creates its database engine at import time. Point this test process
# at a disposable database and isolated credential set before importing the app.
_TEST_ROOT = Path(f"/tmp/confluence-ticketing-pytest-{secrets.token_hex(6)}")
_TEST_ROOT.mkdir(parents=True, exist_ok=True)
os.environ["DATABASE_URL"] = f"sqlite:///{_TEST_ROOT / 'workflow.db'}"
os.environ["DATA_DIR"] = str(_TEST_ROOT / "data")
os.environ["SESSION_SECRET"] = secrets.token_urlsafe(48)
os.environ["QR_TOKEN_PEPPER"] = secrets.token_urlsafe(48)
os.environ["APP_FERNET_KEY"] = base64.urlsafe_b64encode(secrets.token_bytes(32)).decode()
os.environ["BOOTSTRAP_ADMIN_EMAIL"] = "owner@example.test"
os.environ["BOOTSTRAP_ADMIN_PASSWORD"] = "Owner-passphrase-2026!"
os.environ["COOKIE_SECURE"] = "false"

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from PIL import Image, ImageDraw  # noqa: E402
from app.db import Base, SessionLocal, engine  # noqa: E402
from app.main import app  # noqa: E402
from app.models import Attendee, Event, Ticket  # noqa: E402
from app.security import decrypt_qr_token  # noqa: E402


def test_import_ticket_scan_and_attendance_workflow():
    Base.metadata.create_all(bind=engine)

    workbook_buffer = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Name", "Year", "Email", "Phone"])
    sheet.append(["Asha Rao", "FY", "asha@example.test", "9876500001"])
    sheet.append(["Rohan Das", "SY", "rohan@example.test", "9876500002"])
    sheet.append(["Meera Kulkarni", "TY", "meera@example.test", "9876500003"])
    workbook.save(workbook_buffer)
    workbook_bytes = workbook_buffer.getvalue()

    artwork = Image.new("RGB", (1200, 700), "white")
    draw = ImageDraw.Draw(artwork)
    draw.rectangle((0, 0, 1199, 699), outline="#183B56", width=16)
    draw.text((60, 55), "CONFLUENCE SAMPLE TEMPLATE", fill="#183B56")
    artwork_buffer = io.BytesIO()
    artwork.save(artwork_buffer, format="PNG")
    artwork_bytes = artwork_buffer.getvalue()

    with TestClient(app) as client:
        assert client.get("/api/health/ready").status_code == 200
        response = client.post(
            "/api/auth/login",
            json={"email": "owner@example.test", "password": "Owner-passphrase-2026!"},
        )
        assert response.status_code == 200
        csrf = client.get("/api/auth/me").json()["csrfToken"]
        headers = {"X-CSRF-Token": csrf}
        upload = {
            "file": (
                "confluence-students.xlsx",
                workbook_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        }

        response = client.post("/api/imports/inspect", headers=headers, files=upload)
        assert response.status_code == 200
        mapping = response.json()["autoMapping"]
        assert mapping["name"] == "Name"
        assert mapping["email"] == "Email"

        response = client.post(
            "/api/imports/preview",
            headers=headers,
            data={"mapping": json.dumps(mapping)},
            files=upload,
        )
        assert response.status_code == 200
        preview = response.json()
        assert preview["summary"]["valid"] == 3
        response = client.post(f"/api/imports/{preview['previewId']}/commit", headers=headers)
        assert response.status_code == 200
        assert response.json()["added"] == 3

        # Reimporting the same workbook must not create a second set of attendees.
        response = client.post(
            "/api/imports/preview",
            headers=headers,
            data={"mapping": json.dumps(mapping)},
            files=upload,
        )
        assert response.status_code == 200
        duplicate_preview = response.json()
        assert duplicate_preview["summary"]["duplicates"] == 3
        response = client.post(
            f"/api/imports/{duplicate_preview['previewId']}/commit", headers=headers
        )
        assert response.status_code == 200
        assert response.json()["added"] == 0

        response = client.post(
            "/api/template/image",
            headers=headers,
            files={"file": ("sample.png", artwork_bytes, "image/png")},
        )
        assert response.status_code == 200
        placements = response.json()["placements"]
        response = client.put("/api/template/placements", headers=headers, json=placements)
        assert response.status_code == 200

        response = client.post("/api/tickets/generate", headers=headers)
        assert response.status_code == 200
        assert response.headers.get("X-Tickets-Included") == "3"
        archive = zipfile.ZipFile(io.BytesIO(response.content))
        archive_names = archive.namelist()
        ticket_files = [name for name in archive_names if name.endswith(".png")]
        assert len(ticket_files) == 3
        assert "ticket_manifest.csv" in archive_names
        assert Image.open(io.BytesIO(archive.read(ticket_files[0]))).size == (1200, 700)

        with SessionLocal() as db:
            event = db.query(Event).filter_by(slug="confluence").first()
            attendees = db.query(Attendee).filter_by(event_id=event.id).all()
            tickets = db.query(Ticket).filter_by(event_id=event.id).all()
            ticket = next(t for t in tickets if t.attendee_id == attendees[0].id)
            token = decrypt_qr_token(ticket.qr_token_encrypted)
            ticket_id = ticket.id

        # Regeneration must retain the credential already printed on tickets.
        response = client.post("/api/tickets/generate", headers=headers)
        assert response.status_code == 200
        with SessionLocal() as db:
            assert decrypt_qr_token(db.get(Ticket, ticket_id).qr_token_encrypted) == token

        request_id = str(uuid.uuid4())
        scan_payload = {"qrToken": f"CFL1.{token}", "clientScanId": request_id}
        response = client.post("/api/check-ins", headers=headers, json=scan_payload)
        assert response.status_code == 200
        assert response.json()["result"] == "CHECKED_IN"
        # An identical request ID replays the saved result, not another admission.
        response = client.post("/api/check-ins", headers=headers, json=scan_payload)
        assert response.status_code == 200
        assert response.json()["result"] == "CHECKED_IN"
        # A new scan action for the same ticket is rejected as a duplicate.
        response = client.post(
            "/api/check-ins",
            headers=headers,
            json={"qrToken": f"CFL1.{token}", "clientScanId": str(uuid.uuid4())},
        )
        assert response.status_code == 409
        assert response.json()["result"] == "ALREADY_CHECKED_IN"
        response = client.post(
            "/api/check-ins",
            headers=headers,
            json={"qrToken": "CFL1.invalid-long-token-123", "clientScanId": str(uuid.uuid4())},
        )
        assert response.status_code == 404
        assert response.json()["result"] == "INVALID_TICKET"

        report_response = client.get("/api/reports/attendance.xlsx")
        assert report_response.status_code == 200
        report = load_workbook(io.BytesIO(report_response.content), read_only=True, data_only=True).active
        report_rows = list(report.values)
        assert report_rows[0] == (
            "Name", "Year", "Email", "Phone", "Attendance Status"
        )
        assert len(report_rows) == 4
        assert sum(row[-1] == "CHECKED IN" for row in report_rows[1:]) == 1

        other_ticket = next(ticket for ticket in tickets if ticket.id != ticket_id)
        response = client.post(
            f"/api/tickets/{other_ticket.id}/revoke", headers=headers
        )
        assert response.status_code == 200
        with SessionLocal() as db:
            revoked_token = decrypt_qr_token(db.get(Ticket, other_ticket.id).qr_token_encrypted)
        response = client.post(
            "/api/check-ins",
            headers=headers,
            json={"qrToken": f"CFL1.{revoked_token}", "clientScanId": str(uuid.uuid4())},
        )
        assert response.status_code == 410
        assert response.json()["result"] == "TICKET_REVOKED"

        # Changing the owner password keeps the current session but invalidates copied sessions.
        secondary = TestClient(app)
        secondary.cookies.update(client.cookies)
        assert secondary.get("/api/auth/me").json().get("authenticated") is True
        response = client.post(
            "/api/auth/change-password",
            headers=headers,
            json={
                "oldPassword": "Owner-passphrase-2026!",
                "newPassword": "New-owner-passphrase-2026!",
            },
        )
        assert response.status_code == 200
        assert client.get("/api/auth/me").json().get("authenticated") is True
        assert secondary.get("/api/auth/me").json().get("authenticated") is False
        response = client.post(
            "/api/auth/login",
            json={"email": "owner@example.test", "password": "New-owner-passphrase-2026!"},
        )
        assert response.status_code == 200
        secondary.close()

    engine.dispose()
